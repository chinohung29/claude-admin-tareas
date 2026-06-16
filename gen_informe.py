"""
Genera informe diario IVA compras comparando ARCA vs ODOO.
Uso: python gen_informe.py <DD-M> (ej: 1-6)
"""

import sys, os, re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))

TIPO_MAP = {
    "1 - Factura A": "FA-A",
    "2 - Nota de Débito A": "ND-A",
    "3 - Nota de Crédito A": "NC-A",
    "6 - Factura B": "FA-B",
    "11 - Factura C": "FA-C",
    "13 - Nota de Crédito C": "NC-C",
    "Factura A": "FA-A",
    "Nota de Débito A": "ND-A",
    "Nota de Crédito A": "NC-A",
    "Factura B": "FA-B",
    "Factura C": "FA-C",
    "Nota de Crédito C": "NC-C",
}

def normalize_tipo(t):
    if pd.isna(t):
        return ""
    t = str(t).strip()
    for k, v in TIPO_MAP.items():
        if k.lower() in t.lower():
            return v
    return t

def build_arca_key(row, col_pv, col_num):
    tipo = normalize_tipo(row.get("Tipo", ""))
    pv = str(row.get(col_pv, "")).strip().split(".")[0].zfill(5)
    num = str(row.get(col_num, "")).strip().split(".")[0].zfill(8)
    return f"{tipo} {pv}-{num}"

def load_arca(path):
    df = pd.read_excel(path, header=1, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    col_pv = next((c for c in df.columns if "Punto" in c or "punto" in c or "PV" in c), None)
    col_num = next((c for c in df.columns if "mero" in c and "Desde" in c), None)
    if col_pv is None:
        col_pv = [c for c in df.columns if "Venta" in c or "venta" in c]
        col_pv = col_pv[0] if col_pv else df.columns[2]
    if col_num is None:
        col_num = [c for c in df.columns if "mero" in c]
        col_num = col_num[0] if col_num else df.columns[3]
    col_imp = next((c for c in df.columns if "Imp" in c and "Total" in c), None)
    col_neto = next((c for c in df.columns if "Neto" in c), None)
    col_cuit = next((c for c in df.columns if "Emisor" in c and "Doc" not in c and "Denom" not in c and "Nom" not in c), None)
    if col_cuit is None:
        col_cuit = next((c for c in df.columns if "Doc" in c or "CUIT" in c or "cuit" in c), None)
    col_denom = next((c for c in df.columns if "Denom" in c or "Nom" in c and "Emisor" in c), None)
    col_fecha = next((c for c in df.columns if "Fecha" in c or "fecha" in c), None)
    df = df[df["Tipo"].notna() & (df["Tipo"].str.strip() != "")]
    records = []
    for _, row in df.iterrows():
        key = build_arca_key(row, col_pv, col_num)
        imp_total = row.get(col_imp, 0) if col_imp else 0
        try:
            imp_total = float(str(imp_total).replace(",", ".").strip())
        except Exception:
            imp_total = 0.0
        neto = row.get(col_neto, 0) if col_neto else 0
        try:
            neto = float(str(neto).replace(",", ".").strip())
        except Exception:
            neto = 0.0
        cuit = str(row.get(col_cuit, "")).strip() if col_cuit else ""
        denom = str(row.get(col_denom, "")).strip() if col_denom else ""
        fecha = str(row.get(col_fecha, "")).strip() if col_fecha else ""
        tipo_raw = normalize_tipo(row.get("Tipo", ""))
        records.append({
            "key": key, "tipo": tipo_raw, "fecha": fecha,
            "cuit_emisor": cuit, "denominacion": denom,
            "neto": neto, "imp_total_arca": imp_total,
        })
    return pd.DataFrame(records)

def safe_cuit(v):
    try:
        return str(int(float(str(v).strip())))
    except Exception:
        return str(v).strip()

def build_odoo_key(comp_str):
    if pd.isna(comp_str):
        return ""
    s = str(comp_str).strip()
    m = re.match(r"([A-Z]{2}-[A-Z])\s+(\d+)-(\d+)", s)
    if m:
        pv = m.group(2).zfill(5)
        num = m.group(3).zfill(8)
        return f"{m.group(1)} {pv}-{num}"
    for k, v in TIPO_MAP.items():
        if s.lower().startswith(k.lower()):
            rest = s[len(k):].strip()
            m2 = re.match(r"(\d+)-(\d+)", rest)
            if m2:
                pv = m2.group(1).zfill(5)
                num = m2.group(2).zfill(8)
                return f"{v} {pv}-{num}"
    return s

def load_odoo(path):
    df = pd.read_excel(path, header=2, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    comp_col = df.columns[0]
    fecha_col = next((c for c in df.columns if c.lower() == "fecha"), None)
    nombre_col = next((c for c in df.columns if "nombre" in c.lower() or "proveedor" in c.lower()), None)
    cuit_col = next((c for c in df.columns if "cuit" in c.lower()), None)
    total_col = next((c for c in df.columns if c.lower() == "total"), None)
    df = df[df[comp_col].notna()]
    df = df[~df[comp_col].str.strip().str.lower().str.startswith("tota")]
    df = df[df[comp_col].str.strip() != ""]
    records = []
    for _, row in df.iterrows():
        comp_val = str(row[comp_col]).strip()
        key = build_odoo_key(comp_val)
        total = row.get(total_col, 0) if total_col else 0
        try:
            total = float(str(total).replace(",", ".").strip())
        except Exception:
            total = 0.0
        cuit = safe_cuit(row.get(cuit_col, "")) if cuit_col else ""
        nombre = str(row.get(nombre_col, "")).strip() if nombre_col else ""
        fecha = str(row.get(fecha_col, "")).strip() if fecha_col else ""
        records.append({
            "key": key, "comp_raw": comp_val, "fecha_odoo": fecha,
            "nombre": nombre, "cuit": cuit, "total_odoo": total,
        })
    return pd.DataFrame(records)

H_FILL = PatternFill("solid", fgColor="1F4E79")
H_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill("solid", fgColor="D6E4F0")
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
LABEL_FONT = Font(bold=True, size=11)
THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def header_row(ws, row_idx, values):
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=col, value=val)
        c.fill = H_FILL
        c.font = H_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

def data_row(ws, row_idx, values, alt=False):
    fill = ALT_FILL if alt else PatternFill()
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=col, value=val)
        c.fill = fill
        c.border = BORDER
        c.alignment = Alignment(vertical="center")

def auto_width(ws, min_w=12, max_w=50):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, min_w), max_w)

def fmt_money(v):
    try:
        return round(float(v), 2)
    except Exception:
        return v

def main():
    if len(sys.argv) < 2:
        print("Uso: python gen_informe.py <DD-M>")
        sys.exit(1)

    fecha_tag = sys.argv[1]
    arca_path = os.path.join(BASE, "data", "arca", f"arca {fecha_tag}.xlsx")
    odoo_path = os.path.join(BASE, "data", "odoo", f"odoo {fecha_tag}.xlsx")
    if not os.path.exists(odoo_path):
        odoo_path2 = os.path.join(BASE, "data", "odoo", f"Odoo {fecha_tag}.xlsx")
        if os.path.exists(odoo_path2):
            odoo_path = odoo_path2

    informe_dir = os.path.join(BASE, "reports")
    os.makedirs(informe_dir, exist_ok=True)
    out_path = os.path.join(informe_dir, f"informe {fecha_tag}.xlsx")

    if not os.path.exists(arca_path):
        print(f"ADVERTENCIA: No se encontró {arca_path}")
        sys.exit(2)
    if not os.path.exists(odoo_path):
        print(f"ADVERTENCIA: No se encontró {odoo_path}")
        sys.exit(2)

    print(f"Leyendo ARCA: {arca_path}")
    arca = load_arca(arca_path)
    print(f"  {len(arca)} registros en ARCA")

    print(f"Leyendo ODOO: {odoo_path}")
    odoo = load_odoo(odoo_path)
    print(f"  {len(odoo)} registros en ODOO")

    arca_keys = set(arca["key"])
    odoo_keys = set(odoo["key"])

    en_ambos_keys = arca_keys & odoo_keys
    faltantes_keys = arca_keys - odoo_keys
    solo_odoo_keys = odoo_keys - arca_keys

    print(f"En ambos: {len(en_ambos_keys)} | Faltantes en ODOO: {len(faltantes_keys)} | Solo en ODOO: {len(solo_odoo_keys)}")

    merged = arca[arca["key"].isin(en_ambos_keys)].merge(
        odoo[odoo["key"].isin(en_ambos_keys)][["key", "total_odoo"]],
        on="key", how="left"
    )
    merged["diferencia"] = merged["imp_total_arca"] - merged["total_odoo"]
    difs = merged[abs(merged["diferencia"]) > 0.02].copy()

    total_arca = arca["imp_total_arca"].sum()
    total_odoo = odoo["total_odoo"].sum()

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Resumen"
    ws1.freeze_panes = "A1"
    ws1.merge_cells("A1:D1")
    title = ws1["A1"]
    title.value = f"INFORME IVA COMPRAS — {fecha_tag.replace('-', '/')} (fecha {fecha_tag})"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28
    start_row = 3

    labels = [
        ("Cantidad facturas ARCA", len(arca)),
        ("Cantidad facturas ODOO", len(odoo)),
        ("Comprobantes en ambos sistemas", len(en_ambos_keys)),
        ("Faltantes en ODOO (en ARCA, no en ODOO)", len(faltantes_keys)),
        ("Solo en ODOO (no en ARCA)", len(solo_odoo_keys)),
        ("Comprobantes con diferencia de monto", len(difs)),
        ("", ""),
        ("Total Imp. Total ARCA ($)", fmt_money(total_arca)),
        ("Total ODOO ($)", fmt_money(total_odoo)),
        ("Diferencia total ($)", fmt_money(total_arca - total_odoo)),
    ]

    for i, (lbl, val) in enumerate(labels):
        r = start_row + i
        c1 = ws1.cell(row=r, column=1, value=lbl)
        c1.font = LABEL_FONT
        c1.border = BORDER
        c2 = ws1.cell(row=r, column=2, value=val)
        c2.border = BORDER
        c2.alignment = Alignment(horizontal="right")
        if isinstance(val, float):
            c2.number_format = '#,##0.00'

    ws1.column_dimensions["A"].width = 46
    ws1.column_dimensions["B"].width = 22

    ws2 = wb.create_sheet("Faltantes en ODOO")
    faltantes_df = arca[arca["key"].isin(faltantes_keys)].copy()
    h2 = ["Clave", "Tipo", "Fecha", "CUIT Emisor", "Denominación", "Neto Gravado", "Imp. Total ARCA"]
    header_row(ws2, 1, h2)
    ws2.freeze_panes = "A2"
    for i, (_, row) in enumerate(faltantes_df.iterrows()):
        data_row(ws2, i + 2, [
            row["key"], row["tipo"], row["fecha"],
            row["cuit_emisor"], row["denominacion"],
            fmt_money(row["neto"]), fmt_money(row["imp_total_arca"])
        ], alt=(i % 2 == 1))
        ws2.cell(row=i + 2, column=6).number_format = '#,##0.00'
        ws2.cell(row=i + 2, column=7).number_format = '#,##0.00'
    auto_width(ws2)

    ws3 = wb.create_sheet("Diferencias de Monto")
    h3 = ["Clave", "Tipo", "Fecha", "CUIT Emisor", "Denominación", "Imp. Total ARCA", "Total ODOO", "Diferencia"]
    header_row(ws3, 1, h3)
    ws3.freeze_panes = "A2"
    for i, (_, row) in enumerate(difs.iterrows()):
        data_row(ws3, i + 2, [
            row["key"], row["tipo"], row["fecha"],
            row["cuit_emisor"], row["denominacion"],
            fmt_money(row["imp_total_arca"]),
            fmt_money(row["total_odoo"]),
            fmt_money(row["diferencia"])
        ], alt=(i % 2 == 1))
        for col_idx in [6, 7, 8]:
            ws3.cell(row=i + 2, column=col_idx).number_format = '#,##0.00'
    auto_width(ws3)

    wb.save(out_path)
    print(f"\nInforme generado: {out_path}")
    print(f"  Hoja 'Resumen': estadísticas generales")
    print(f"  Hoja 'Faltantes en ODOO': {len(faltantes_df)} comprobantes")
    print(f"  Hoja 'Diferencias de Monto': {len(difs)} comprobantes")

if __name__ == "__main__":
    main()
