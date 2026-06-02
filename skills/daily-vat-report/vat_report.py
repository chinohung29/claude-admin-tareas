#!/usr/bin/env python3
"""Daily VAT reconciliation report: Odoo vs ARCA."""

import re
import sys
import argparse
from pathlib import Path
from datetime import datetime, date

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


ARCA_TYPE_MAP = {
    "1 - Factura A": "FA-A",
    "2 - Nota de Débito A": "ND-A",
    "3 - Nota de Crédito A": "NC-A",
    "4 - Recibo A": "RE-A",
    "6 - Factura B": "FA-B",
    "7 - Nota de Débito B": "ND-B",
    "8 - Nota de Crédito B": "NC-B",
    "11 - Factura C": "FA-C",
    "12 - Nota de Débito C": "ND-C",
    "13 - Nota de Crédito C": "NC-C",
    "51 - Factura M": "FA-M",
    "201 - Factura de Crédito Electrónica MiPyMEs (FCE) A": "FCE-A",
}

ODOO_TYPE_RE = re.compile(
    r"^(FA|NC|ND|RE|OC|FCE)-([A-Z])\s+(\d{5})-(\d{8})$"
)


def parse_odoo(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, header=None)
    header_row = None
    for i, row in df.iterrows():
        vals = [str(v).strip() for v in row.values if pd.notna(v)]
        if "Fecha" in vals and "CUIT" in vals:
            header_row = i
            break
    if header_row is None:
        raise ValueError(f"No se encontró la fila de encabezado en {path}")

    df = pd.read_excel(path, header=header_row)
    df.columns = [str(c).strip() for c in df.columns]
    first_col = df.columns[0]
    df = df[df[first_col].notna()].copy()
    df = df[~df[first_col].astype(str).str.startswith("Total")].copy()
    df.rename(columns={first_col: "Comprobante"}, inplace=True)
    df["Comprobante"] = df["Comprobante"].astype(str).str.strip()

    records = []
    for _, row in df.iterrows():
        comp = row["Comprobante"]
        m = ODOO_TYPE_RE.match(comp)
        if not m:
            continue
        tipo_code, letra, pto_vta, numero = m.groups()
        key = f"{tipo_code}-{letra} {pto_vta}-{numero}"
        total = _to_float(row.get("Total", 0))
        gravado = _to_float(row.get("Gravado", 0))
        iva_21 = _to_float(row.get("IVA 21%", 0))
        cuit = str(row.get("CUIT", "")).replace("-", "").strip()
        nombre = str(row.get("Nombre", "")).strip()
        fecha = row.get("Fecha", "")
        records.append({
            "key": key,
            "comprobante": comp,
            "fecha": fecha,
            "nombre": nombre,
            "cuit": cuit,
            "gravado": gravado,
            "iva_21": iva_21,
            "total": total,
            "source": "odoo",
        })
    return pd.DataFrame(records)


def parse_arca(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, header=None)
    header_row = None
    for i, row in df.iterrows():
        vals = [str(v).strip() for v in row.values if pd.notna(v)]
        if "Fecha" in vals and "Tipo" in vals:
            header_row = i
            break
    if header_row is None:
        raise ValueError(f"No se encontró la fila de encabezado en {path}")

    df = pd.read_excel(path, header=header_row)
    df.columns = [str(c).strip() for c in df.columns]
    df = df[df["Fecha"].notna()].copy()

    records = []
    for _, row in df.iterrows():
        tipo_str = str(row.get("Tipo", "")).strip()
        tipo_code = ARCA_TYPE_MAP.get(tipo_str)
        if not tipo_code:
            continue
        pto_vta = int(row.get("Punto de Venta", 0))
        numero = int(row.get("Número Desde", 0))
        key = f"{tipo_code} {pto_vta:05d}-{numero:08d}"
        total = _to_float(row.get("Imp. Total", 0))
        gravado = _to_float(row.get("Neto Gravado Total", 0))
        iva_21 = _to_float(row.get("IVA 21%", 0))
        cuit = str(row.get("Nro. Doc. Emisor", "")).replace("-", "").strip()
        nombre = str(row.get("Denominación Emisor", "")).strip()
        fecha = row.get("Fecha", "")
        records.append({
            "key": key,
            "comprobante": key,
            "fecha": fecha,
            "nombre": nombre,
            "cuit": cuit,
            "gravado": gravado,
            "iva_21": iva_21,
            "total": total,
            "source": "arca",
        })
    return pd.DataFrame(records)


def _to_float(val) -> float:
    if pd.isna(val) or val == "" or val is None:
        return 0.0
    try:
        return float(str(val).replace(",", ""))
    except (ValueError, TypeError):
        return 0.0


def reconcile(odoo_df: pd.DataFrame, arca_df: pd.DataFrame) -> dict:
    odoo_keys = set(odoo_df["key"])
    arca_keys = set(arca_df["key"])

    only_arca = arca_keys - odoo_keys
    only_odoo = odoo_keys - arca_keys
    common = odoo_keys & arca_keys

    only_arca_detail = arca_df[arca_df["key"].isin(only_arca)].copy()
    only_odoo_detail = odoo_df[odoo_df["key"].isin(only_odoo)].copy()

    diffs = []
    for key in common:
        odoo_row = odoo_df[odoo_df["key"] == key].iloc[0]
        arca_row = arca_df[arca_df["key"] == key].iloc[0]
        total_diff = round(odoo_row["total"] - arca_row["total"], 2)
        if abs(total_diff) > 0.01:
            diffs.append({
                "comprobante": key,
                "nombre": odoo_row["nombre"],
                "cuit": odoo_row["cuit"],
                "total_odoo": odoo_row["total"],
                "total_arca": arca_row["total"],
                "diferencia": total_diff,
            })

    return {
        "total_arca": len(arca_df),
        "total_odoo": len(odoo_df),
        "only_arca_count": len(only_arca),
        "only_odoo_count": len(only_odoo),
        "common_count": len(common),
        "diff_count": len(diffs),
        "only_arca_detail": only_arca_detail,
        "only_odoo_detail": only_odoo_detail,
        "amount_diffs": pd.DataFrame(diffs) if diffs else pd.DataFrame(),
    }


def generate_report(result: dict, output_path: str, report_date: str):
    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    subheader_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws = wb.active
    ws.title = "Resumen"

    ws.merge_cells("A1:D1")
    ws["A1"] = f"Informe de Conciliación IVA — {report_date}"
    ws["A1"].font = Font(bold=True, size=14, color="2F5496")
    ws["A1"].alignment = Alignment(horizontal="center")

    summary_data = [
        ("Comprobantes en ARCA", result["total_arca"]),
        ("Comprobantes en Odoo", result["total_odoo"]),
        ("Comprobantes coincidentes", result["common_count"]),
        ("Solo en ARCA (faltan en Odoo)", result["only_arca_count"]),
        ("Solo en Odoo (faltan en ARCA)", result["only_odoo_count"]),
        ("Con diferencia de importe", result["diff_count"]),
    ]

    for i, (label, value) in enumerate(summary_data, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=value)
        cell.alignment = Alignment(horizontal="center")
        if "faltan" in label.lower() and value > 0:
            cell.font = Font(bold=True, color="FF0000")

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 15

    if not result["only_arca_detail"].empty:
        ws2 = wb.create_sheet("Faltan en Odoo")
        _write_detail_sheet(ws2, result["only_arca_detail"],
                           "Comprobantes en ARCA que faltan en Odoo",
                           header_font, header_fill, border)

    if not result["only_odoo_detail"].empty:
        ws3 = wb.create_sheet("Faltan en ARCA")
        _write_detail_sheet(ws3, result["only_odoo_detail"],
                           "Comprobantes en Odoo que faltan en ARCA",
                           header_font, header_fill, border)

    if not result["amount_diffs"].empty:
        ws4 = wb.create_sheet("Diferencias de Importe")
        _write_diff_sheet(ws4, result["amount_diffs"],
                         header_font, header_fill, border)

    wb.save(output_path)
    return output_path


def _write_detail_sheet(ws, df, title, header_font, header_fill, border):
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12, color="2F5496")

    cols = ["comprobante", "fecha", "nombre", "cuit", "total"]
    headers = ["Comprobante", "Fecha", "Nombre", "CUIT", "Total"]
    widths = [25, 15, 40, 18, 18]

    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + j)].width = w

    for i, (_, row) in enumerate(df.iterrows(), start=4):
        for j, col in enumerate(cols, start=1):
            val = row[col]
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = border
            if col == "total":
                cell.number_format = '#,##0.00'

    total_row = 4 + len(df)
    ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=df["total"].sum()).number_format = '#,##0.00'
    ws.cell(row=total_row, column=5).font = Font(bold=True)


def _write_diff_sheet(ws, df, header_font, header_fill, border):
    ws.merge_cells("A1:F1")
    ws["A1"] = "Diferencias de Importe entre Odoo y ARCA"
    ws["A1"].font = Font(bold=True, size=12, color="2F5496")

    headers = ["Comprobante", "Nombre", "CUIT", "Total Odoo", "Total ARCA", "Diferencia"]
    cols = ["comprobante", "nombre", "cuit", "total_odoo", "total_arca", "diferencia"]
    widths = [25, 40, 18, 18, 18, 18]

    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + j)].width = w

    for i, (_, row) in enumerate(df.iterrows(), start=4):
        for j, col in enumerate(cols, start=1):
            val = row[col]
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = border
            if col in ("total_odoo", "total_arca", "diferencia"):
                cell.number_format = '#,##0.00'
                if col == "diferencia" and val != 0:
                    cell.font = Font(color="FF0000")


def main():
    parser = argparse.ArgumentParser(description="Generar informe de conciliación IVA")
    parser.add_argument("odoo_file", help="Archivo Excel de Odoo")
    parser.add_argument("arca_file", help="Archivo Excel de ARCA")
    parser.add_argument("-o", "--output", help="Archivo de salida (default: reports/informe_iva_FECHA.xlsx)")
    parser.add_argument("-d", "--date", help="Fecha del informe (YYYY-MM-DD)", default=str(date.today()))
    args = parser.parse_args()

    if not args.output:
        reports_dir = Path(__file__).resolve().parent.parent.parent / "reports"
        reports_dir.mkdir(exist_ok=True)
        args.output = str(reports_dir / f"informe_iva_{args.date}.xlsx")

    print(f"Procesando Odoo: {args.odoo_file}")
    odoo_df = parse_odoo(args.odoo_file)
    print(f"  → {len(odoo_df)} comprobantes")

    print(f"Procesando ARCA: {args.arca_file}")
    arca_df = parse_arca(args.arca_file)
    print(f"  → {len(arca_df)} comprobantes")

    print("Conciliando...")
    result = reconcile(odoo_df, arca_df)

    print(f"\n--- Resumen ---")
    print(f"Comprobantes en ARCA:           {result['total_arca']}")
    print(f"Comprobantes en Odoo:           {result['total_odoo']}")
    print(f"Coincidentes:                   {result['common_count']}")
    print(f"Solo en ARCA (faltan en Odoo):  {result['only_arca_count']}")
    print(f"Solo en Odoo (faltan en ARCA):  {result['only_odoo_count']}")
    print(f"Con diferencia de importe:      {result['diff_count']}")

    output = generate_report(result, args.output, args.date)
    print(f"\nInforme generado: {output}")


if __name__ == "__main__":
    main()
